from agent2db.documents import Extracted, chunk_text, extract_text, kind_of


def test_kind_detection():
    assert kind_of("report.PDF") == "pdf"
    assert kind_of("notes.md") == "markdown"
    assert kind_of("spec.docx") == "docx"
    assert kind_of("page.html") == "html"
    assert kind_of("data.csv") == "text"
    assert kind_of("blob", "text/markdown") == "markdown"
    try:
        kind_of("photo.png")
    except ValueError as exc:
        assert "Unsupported" in str(exc)
    else:
        raise AssertionError("png should be rejected")


def test_markdown_chunks_carry_heading_paths():
    text = "# Billing\n\nIntro paragraph.\n\n## Refunds\n\nRefunds take 5 days.\n\n# Shipping\n\nShips in 2 days."
    chunks = chunk_text(Extracted(text, "markdown"), max_chars=1200, min_chars=0)
    assert [c.heading for c in chunks] == ["Billing", "Billing > Refunds", "Shipping"]
    assert chunks[1].content == "Refunds take 5 days."
    assert [c.idx for c in chunks] == [0, 1, 2]


def test_long_text_is_packed_with_overlap():
    paragraphs = [f"Paragraph {i}. " + ("word " * 60).strip() for i in range(12)]
    chunks = chunk_text("\n\n".join(paragraphs), max_chars=800, overlap=100, min_chars=100)
    assert len(chunks) > 3
    assert all(c.char_count <= 800 + 100 for c in chunks)
    # Overlap: the start of chunk 2 repeats the tail of chunk 1.
    tail = chunks[0].content[-40:]
    assert tail in chunks[1].content


def test_oversized_paragraph_is_split_at_sentences():
    text = ". ".join(f"Sentence number {i} is here" for i in range(80)) + "."
    chunks = chunk_text(text, max_chars=500, overlap=0, min_chars=0)
    assert all(c.char_count <= 500 for c in chunks)
    assert "".join(c.content for c in chunks).count("Sentence number 79") == 1


def test_tiny_trailing_chunk_is_merged():
    text = ("x" * 700) + "\n\n" + ("y" * 700) + "\n\nshort tail"
    chunks = chunk_text(text, max_chars=800, overlap=0, min_chars=200)
    assert chunks[-1].content.endswith("short tail")
    assert len(chunks) == 2


def test_html_and_text_extraction():
    html = b"<html><head><style>p{}</style></head><body><h1>Title</h1><p>Hello &amp; welcome.</p><script>x()</script></body></html>"
    out = extract_text("page.html", html)
    assert out.kind == "html" and "# Title" in out.text and "Hello & welcome." in out.text and "x()" not in out.text
    out = extract_text("a.txt", b"line1\r\nline2")
    assert out.text == "line1\nline2" and out.kind == "text"


def test_pdf_pages_are_tracked():
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(2):
        writer.add_blank_page(width=200, height=200)
    import io

    buf = io.BytesIO()
    writer.write(buf)
    out = extract_text("blank.pdf", buf.getvalue())
    assert out.kind == "pdf" and out.pages == 2 and len(out.page_starts) == 2


def test_docx_headings_become_markdown_headings():
    import io

    import docx

    document = docx.Document()
    document.add_heading("Policy", level=1)
    document.add_paragraph("Returns are accepted within 30 days.")
    document.add_heading("Exceptions", level=2)
    document.add_paragraph("Final sale items are excluded.")
    buf = io.BytesIO()
    document.save(buf)
    out = extract_text("policy.docx", buf.getvalue())
    assert out.kind == "docx"
    chunks = chunk_text(out, min_chars=0)
    assert [c.heading for c in chunks] == ["Policy", "Policy > Exceptions"]
